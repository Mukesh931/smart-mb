"""
Smart-MB :: Report Engine
Generates the official-style Form No. 23 Measurement Book (MB) as PDF and Excel,
always using the clean legal description text from the Master CSR database.

Deviations: Total Measured vs Tendered (Estimate) quantity, per item, with
Excess / Saving amount computed at Master CSR rates.
"""
from __future__ import annotations

import io
from datetime import datetime

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (KeepTogether, PageBreak, Paragraph, SimpleDocTemplate,
                                Spacer, Table, TableStyle)

INK = colors.HexColor("#0f172a")
MUTED = colors.HexColor("#475569")
LINE = colors.HexColor("#94a3b8")
BAND = colors.HexColor("#e2e8f0")
ACCENT = colors.HexColor("#0b5c8a")

PAGE = landscape(A4)
MARGIN = 12 * mm


# ------------------------------------------------------------------ formatting
def fmt_num(value: float | None, dp: int = 2) -> str:
    if value is None:
        return ""
    if abs(value - round(value)) < 1e-9 and dp <= 2:
        value = round(value)
        s = f"{int(value):,}"
        return _indian(s)
    return _indian(f"{value:,.{dp}f}")


def _indian(s: str) -> str:
    """1234567.89 -> 12,34,567.89"""
    neg = s.startswith("-")
    s = s.lstrip("-")
    if "." in s:
        whole, frac = s.split(".")
    else:
        whole, frac = s, ""
    whole = whole.replace(",", "")
    if len(whole) > 3:
        head, tail = whole[:-3], whole[-3:]
        parts = []
        while len(head) > 2:
            parts.insert(0, head[-2:])
            head = head[:-2]
        if head:
            parts.insert(0, head)
        whole = ",".join(parts + [tail])
    out = whole + ("." + frac if frac else "")
    return ("-" if neg else "") + out


def inr(value: float | None) -> str:
    return "Rs. " + fmt_num(value, 2) if value not in (None, "") else "-"


def fmt_date(value: str | None) -> str:
    if not value:
        return "____________"
    for f in ("%Y-%m-%d", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M:%S%z"):
        try:
            return datetime.strptime(str(value)[:19], f).strftime("%d/%m/%Y")
        except Exception:
            continue
    return str(value)


# --------------------------------------------------------------------- styles
def _styles():
    ss = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("t", parent=ss["Title"], fontSize=13.5, leading=16, textColor=INK, spaceAfter=1),
        "sub": ParagraphStyle("s", parent=ss["Normal"], fontSize=8.4, leading=11, textColor=MUTED, alignment=TA_CENTER),
        "cell": ParagraphStyle("c", parent=ss["Normal"], fontSize=7.1, leading=8.7, textColor=INK),
        "cellb": ParagraphStyle("cb", parent=ss["Normal"], fontSize=7.1, leading=8.7, textColor=INK,
                                fontName="Helvetica-Bold"),
        "cellr": ParagraphStyle("cr", parent=ss["Normal"], fontSize=7.1, leading=8.7, textColor=INK,
                                alignment=2),
        "h": ParagraphStyle("h", parent=ss["Normal"], fontSize=8.6, leading=11, textColor=colors.white,
                            fontName="Helvetica-Bold", alignment=TA_CENTER),
        "small": ParagraphStyle("sm", parent=ss["Normal"], fontSize=7.6, leading=9.6, textColor=INK),
        "sect": ParagraphStyle("sc", parent=ss["Normal"], fontSize=9.6, leading=12, textColor=ACCENT,
                               fontName="Helvetica-Bold", spaceBefore=6, spaceAfter=3),
    }


def _kv_table(rows: list[tuple[str, str]], styles, widths=(34 * mm, 96 * mm, 34 * mm, 96 * mm)):
    data = []
    for i in range(0, len(rows), 2):
        pair = rows[i:i + 2]
        line = []
        for k, v in pair:
            line += [Paragraph(f"<b>{k}</b>", styles["small"]), Paragraph(str(v or "-"), styles["small"])]
        if len(pair) == 1:
            line += ["", ""]
        data.append(line)
    t = Table(data, colWidths=list(widths), hAlign="LEFT")
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, LINE),
        ("BACKGROUND", (0, 0), (0, -1), BAND),
        ("BACKGROUND", (2, 0), (2, -1), BAND),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 2.6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.6),
    ]))
    return t


# --------------------------------------------------------------- Form 23 (PDF)
def form23_pdf(project: dict, items: list[dict], profile: dict | None = None,
               *, include_deviations: bool = True, measured_only: bool = False) -> bytes:
    """items = [{project_item fields..., 'measurements': [...], 'measured_qty': x}]"""
    styles = _styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=PAGE, leftMargin=MARGIN, rightMargin=MARGIN,
        topMargin=20 * mm, bottomMargin=16 * mm,
        title=f"Form 23 MB {project.get('mb_no') or ''} - {project.get('name')}",
        author="Smart-MB | Maharashtra PWD Electrical",
    )
    width = PAGE[0] - 2 * MARGIN
    flow = []

    def header_footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica-Bold", 10.5)
        canvas.setFillColor(INK)
        canvas.drawCentredString(PAGE[0] / 2, PAGE[1] - 11 * mm, "FORM No. 23  -  MEASUREMENT BOOK")
        canvas.setFont("Helvetica", 7.4)
        canvas.setFillColor(MUTED)
        canvas.drawCentredString(PAGE[0] / 2, PAGE[1] - 15 * mm,
                                 "GOVERNMENT OF MAHARASHTRA  |  PUBLIC WORKS DEPARTMENT (ELECTRICAL WING)")
        canvas.drawString(MARGIN, PAGE[1] - 15 * mm, f"MB No. {project.get('mb_no') or '-'}")
        canvas.drawRightString(PAGE[0] - MARGIN, PAGE[1] - 15 * mm,
                               f"CSR {project.get('csr_fy') or '-'} / {project.get('csr_region') or '-'}")
        canvas.setStrokeColor(LINE)
        canvas.setLineWidth(0.6)
        canvas.line(MARGIN, PAGE[1] - 17 * mm, PAGE[0] - MARGIN, PAGE[1] - 17 * mm)
        canvas.line(MARGIN, 12 * mm, PAGE[0] - MARGIN, 12 * mm)
        canvas.setFont("Helvetica", 7)
        canvas.drawString(MARGIN, 8.6 * mm, "Generated by Smart-MB | Smart Schedule-Mapped Measurement Book system")
        canvas.drawCentredString(PAGE[0] / 2, 8.6 * mm,
                                 "Descriptions & rates as per Master CSR Database (admin-controlled) ")
        canvas.drawRightString(PAGE[0] - MARGIN, 8.6 * mm, f"Page {doc_.page}")
        canvas.restoreState()

    # ---- work particulars
    flow.append(Paragraph((project.get("name") or "").upper(), styles["title"]))
    flow.append(Paragraph(
        f"{project.get('project_code','')}  |  Division: {project.get('division','')}  |  Circle: {project.get('circle','')}",
        styles["sub"]))
    flow.append(Spacer(1, 4 * mm))
    particulars = [
        ("Name of Work", project.get("name", "-")),
        ("Project Code", project.get("project_code", "-")),
        ("Scheme / Head", project.get("scheme", "-")),
        ("Name of Agency", project.get("agency", "-")),
        ("Estimate No.", project.get("estimate_no", "-")),
        ("Technical Sanction No. / Date", f"{project.get('ts_no','-')}  dt. {fmt_date(project.get('ts_date'))}"),
        ("Agreement No.", project.get("agreement_no", "-")),
        ("MB No. / Financial Year", f"{project.get('mb_no','-')} / {project.get('csr_fy','-')}"),
        ("Tendered Amount (Estimate)", inr(project.get("ts_amount"))),
        ("CSR Applicable", f"Electrical CSR {project.get('csr_fy','-')} - {project.get('csr_region','-')} region"),
        ("Measurement Period", profile.get("period", "-") if profile else "-"),
        ("Engineer in Charge", (profile or {}).get("engineer", "-")),
    ]
    flow.append(_kv_table(particulars, styles, widths=(40 * mm, 92 * mm, 40 * mm, 88 * mm)))
    flow.append(Spacer(1, 5 * mm))

    # ---- measurement table
    col_w = [16 * mm, width - (16 + 12 + 17 + 17 + 17 + 22 + 22 + 20 + 16) * mm,
             12 * mm, 17 * mm, 17 * mm, 17 * mm, 22 * mm, 22 * mm, 20 * mm, 16 * mm]
    head = ["Item No.", "Description of Item (as per Master CSR) / Location", "Unit", "No.", "Length (m)",
            "Breadth (m)", "Height / Depth (m)", "Quantity", "Rate (Rs.)", "Amount (Rs.)"]
    data = [[Paragraph(h, styles["h"]) for h in head]]
    grand_qty, grand_amt = 0.0, 0.0
    dev_rows = []

    for idx, it in enumerate(items, start=1):
        meas = it.get("measurements", [])
        measured = it.get("measured_qty") or 0.0
        rate = it.get("rate") or 0.0
        amount = round(measured * rate, 2)
        if measured_only and measured <= 0:
            continue
        grand_qty += measured
        grand_amt += amount

        # item head (serial + master description + CSR code)
        code = it.get("item_code", "")
        head_html = (f"<b>Item {idx}</b> &nbsp;|&nbsp; <b>CSR Code {code}</b> &nbsp;|&nbsp; {it.get('unit','')}<br/>"
                     f"{it.get('description','')}")
        if it.get("is_non_schedule"):
            head_html += f"<br/><i>NON-SCHEDULE ITEM</i> - {it.get('ns_reason') or 'approved extra item, rate as per sanction'}"
        data.append([
            Paragraph(f"<b>{idx}</b><br/>CSR<br/>{code}", styles["cell"]),
            Paragraph(head_html, styles["cell"]),
            Paragraph(it.get("unit", ""), styles["cell"]),
            "", "", "", "", "", "", "",
        ])
        if it.get("tendered_qty") is not None:
            tq = it.get("tendered_qty") or 0
            dev = measured - tq
            pct = (dev / tq * 100) if tq else 0.0
            dev_rows.append({
                "item_code": code, "description": it.get("short_desc") or it.get("description", "")[:90],
                "unit": it.get("unit"), "rate": rate, "tendered": tq, "measured": measured,
                "dev_qty": dev, "dev_pct": pct, "dev_amount": round(dev * rate, 2),
            })

        if not meas:
            data.append(["", Paragraph("<i>Not yet measured / nil</i>", styles["cell"]), "", "", "", "", "", "", "", ""])
        for m in meas:
            label = f"{m.get('room_name') or 'Site'}"
            if m.get("floor"):
                label = f"{m['floor']} - {label}"
            note = m.get("notes") or ""
            loc = f"<b>{label}</b>" + (f"<br/><font size=6 color='#64748b'>{note}</font>" if note else "")
            data.append([
                Paragraph(f"{idx}.{m.get('sub','')}", styles["cell"]),
                Paragraph(loc, styles["cell"]),
                Paragraph(it.get("unit", ""), styles["cell"]),
                Paragraph(fmt_num(m.get("nos"), 0), styles["cellr"]),
                Paragraph(fmt_num(m.get("length"), 2), styles["cellr"]),
                Paragraph(fmt_num(m.get("breadth"), 2) if m.get("breadth") else "", styles["cellr"]),
                Paragraph(fmt_num(m.get("height"), 2) if m.get("height") else "", styles["cellr"]),
                Paragraph(fmt_num(m.get("measured_qty"), 2), styles["cellr"]),
                Paragraph(fmt_num(rate, 2), styles["cellr"]),
                Paragraph(fmt_num(m.get("amount", (m.get("measured_qty") or 0) * rate), 2), styles["cellr"]),
            ])
        data.append([
            "", Paragraph("<b>Total for Item</b>", styles["cellr"]), Paragraph(it.get("unit", ""), styles["cell"]),
            "", "", "", "", Paragraph(f"<b>{fmt_num(measured, 2)}</b>", styles["cellr"]),
            Paragraph(fmt_num(rate, 2), styles["cellr"]), Paragraph(f"<b>{fmt_num(amount, 2)}</b>", styles["cellr"]),
        ])

    # label goes in the first cell of the merged block (a SPAN keeps the leader cell's content)
    data.append([Paragraph("<b>GRAND TOTAL</b>", styles["cellr"]), "", "", "", "", "", "",
                 Paragraph(f"<b>{fmt_num(grand_qty, 2)}</b>", styles["cellr"]), "",
                 Paragraph(f"<b>{fmt_num(grand_amt, 2)}</b>", styles["cellr"])])

    tbl = Table(data, colWidths=col_w, repeatRows=1, hAlign="LEFT")
    tbl.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, LINE),
        ("BACKGROUND", (0, 0), (-1, 0), ACCENT),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 2.4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.4),
        ("BACKGROUND", (0, -1), (-1, -1), BAND),
        ("SPAN", (0, -1), (6, -1)),
    ]))
    flow.append(tbl)
    flow.append(Spacer(1, 4 * mm))
    flow.append(Paragraph(
        f"Total measured value: <b>{inr(grand_amt)}</b> &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"Tendered (estimate) value: <b>{inr(project.get('ts_amount'))}</b> &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"Net deviation: <b>{inr(round(grand_amt - (project.get('ts_amount') or 0), 2))}</b>",
        styles["small"]))
    flow.append(Spacer(1, 5 * mm))

    # ---- certification + signatures
    flow.append(Paragraph(
        "Certified that the above measurements were jointly recorded at site in the presence of the representative of "
        "the contractor, that the quantities shown are as actually executed and that the descriptions, units and rates "
        "adopted are as per the current Schedule of Rates (Electrical CSR) applicable to this work.", styles["small"]))
    flow.append(Spacer(1, 7 * mm))
    sign = Table([
        [Paragraph("<b>Measured by</b><br/><br/><br/><br/>Junior Engineer (Electrical)<br/>PWD Electrical Sub-Division",
                   styles["small"]),
         Paragraph("<b>Checked by</b><br/><br/><br/><br/>Deputy Engineer (Electrical)<br/>PWD Electrical Sub-Division",
                   styles["small"]),
         Paragraph("<b>Contractor / Representative</b><br/><br/><br/><br/>Signature with date<br/>(Joint Measurement)",
                   styles["small"]),
         Paragraph("<b>Countersigned</b><br/><br/><br/><br/>Executive Engineer (Electrical)<br/>PWD Electrical Division",
                   styles["small"])],
    ], colWidths=[width / 4] * 4, hAlign="LEFT")
    sign.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.4, LINE),
                              ("INNERGRID", (0, 0), (-1, -1), 0.4, LINE),
                              ("VALIGN", (0, 0), (-1, -1), "TOP"),
                              ("TOPPADDING", (0, 0), (-1, -1), 5),
                              ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    flow.append(sign)

    # ---- deviation schedule
    if include_deviations and dev_rows:
        flow.append(PageBreak())
        flow.append(Paragraph("DEVIATION STATEMENT  (Measured vs Tendered)", styles["sect"]))
        flow.append(Paragraph(
            "Quantities measured at site compared with the tendered quantities of the Technical Sanction estimate. "
            "Excess / savings are valued at the Master CSR rates to identify items requiring a variation order or "
            "excess-item sanction.", styles["small"]))
        flow.append(Spacer(1, 3 * mm))
        head2 = [Paragraph(h, styles["h"]) for h in
                 ["CSR Code", "Item (short description)", "Unit", "Rate (Rs.)", "Tendered Qty",
                  "Measured Qty", "Deviation", "Dev. %", "Excess (+) / Saving (-) Amount (Rs.)"]]
        d = [head2]
        for r in sorted(dev_rows, key=lambda x: -abs(x["dev_amount"])):
            flag = "EXCESS" if r["dev_qty"] > 0.001 else ("SAVING" if r["dev_qty"] < -0.001 else "-")
            d.append([
                Paragraph(r["item_code"], styles["cell"]),
                Paragraph((r["description"] or "") + (f" <font size=6 color='#b91c1c'>[{flag}]</font>"
                                                      if flag != "-" and abs(r["dev_pct"]) >= 10 else ""), styles["cell"]),
                Paragraph(r["unit"] or "", styles["cell"]),
                Paragraph(fmt_num(r["rate"], 2), styles["cellr"]),
                Paragraph(fmt_num(r["tendered"], 2), styles["cellr"]),
                Paragraph(fmt_num(r["measured"], 2), styles["cellr"]),
                Paragraph(f"{'+' if r['dev_qty'] > 0 else ''}{fmt_num(r['dev_qty'], 2)}", styles["cellr"]),
                Paragraph(f"{r['dev_pct']:+.1f}%", styles["cellr"]),
                Paragraph(fmt_num(r["dev_amount"], 2), styles["cellr"]),
            ])
        tot_t = sum(r["tendered"] for r in dev_rows)
        tot_m = sum(r["measured"] for r in dev_rows)
        tot_a = sum(r["dev_amount"] for r in dev_rows)
        exc = sum(r["dev_amount"] for r in dev_rows if r["dev_amount"] > 0)
        sav = sum(r["dev_amount"] for r in dev_rows if r["dev_amount"] < 0)
        d.append(["", Paragraph("<b>TOTAL</b>", styles["cellr"]), "", "", Paragraph(f"<b>{fmt_num(tot_t, 2)}</b>", styles["cellr"]),
                  Paragraph(f"<b>{fmt_num(tot_m, 2)}</b>", styles["cellr"]), "", "",
                  Paragraph(f"<b>{fmt_num(tot_a, 2)}</b>", styles["cellr"])])
        dt = Table(d, colWidths=[22 * mm, width - 200 * mm, 16 * mm, 22 * mm, 24 * mm, 24 * mm, 22 * mm, 18 * mm, 40 * mm],
                   repeatRows=1, hAlign="LEFT")
        dt.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.4, LINE),
            ("BACKGROUND", (0, 0), (-1, 0), ACCENT),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BACKGROUND", (0, -1), (-1, -1), BAND),
            ("TOPPADDING", (0, 0), (-1, -1), 2.4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.4),
        ]))
        flow.append(dt)
        flow.append(Spacer(1, 4 * mm))
        flow.append(Paragraph(
            f"<b>Excess value: {inr(round(exc,2))}</b> &nbsp;&nbsp;|&nbsp;&nbsp; "
            f"<b>Savings value: {inr(round(sav,2))}</b> &nbsp;&nbsp;|&nbsp;&nbsp; "
            f"<b>Net: {inr(round(exc + sav, 2))}</b><br/>"
            "Note: Items showing excess beyond the permissible variation limit require a formal variation / "
            "excess item approval before recording in the final bill.", styles["small"]))
        flow.append(Spacer(1, 8 * mm))
        flow.append(_kv_table([
            ("Prepared by (JE)", (profile or {}).get("engineer", "____________________")),
            ("Date of preparation", datetime.now().strftime("%d/%m/%Y")),
        ], styles, widths=(40 * mm, 92 * mm, 40 * mm, 88 * mm)))

    doc.build(flow, onFirstPage=header_footer, onLaterPages=header_footer)
    return buf.getvalue()


# ------------------------------------------------------------- Form 23 (Excel)
def form23_xlsx(project: dict, items: list[dict], profile: dict | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    thin = Side(style="thin", color="94A3B8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill("solid", fgColor="0B5C8A")
    band = PatternFill("solid", fgColor="E2E8F0")
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Form-23 Measurements"
    ws["A1"] = "FORM No. 23 - MEASUREMENT BOOK"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "GOVERNMENT OF MAHARASHTRA | PUBLIC WORKS DEPARTMENT (ELECTRICAL WING) | Smart-MB"
    ws["A3"] = f"Name of Work: {project.get('name','')}"
    ws["A4"] = (f"Project Code: {project.get('project_code','')}   |   Estimate No: {project.get('estimate_no','')}"
                f"   |   TS No: {project.get('ts_no','')} dt. {fmt_date(project.get('ts_date'))}")
    ws["A5"] = (f"MB No: {project.get('mb_no','')}   |   Agreement: {project.get('agreement_no','')}"
                f"   |   Agency: {project.get('agency','')}   |   CSR: {project.get('csr_fy','')} / {project.get('csr_region','')}")
    ws["A6"] = (f"Tendered (estimate) amount: {inr(project.get('ts_amount'))}   |   "
                f"Prepared: {datetime.now().strftime('%d/%m/%Y %H:%M')}")

    headers = ["Item No.", "CSR Code", "Description of Item (Master CSR)", "Location / Room", "Floor", "Unit", "No.",
               "Length", "Breadth", "Height", "Quantity", "Rate", "Amount", "Measured On", "Status"]
    ws.append([])
    ws.append(headers)
    hrow = ws.max_row
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=hrow, column=c)
        cell.fill = hdr_fill
        cell.font = Font(bold=True, color="FFFFFF", size=9)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.freeze_panes = ws.cell(row=hrow + 1, column=1)

    total = 0.0
    for idx, it in enumerate(items, start=1):
        rate = it.get("rate") or 0
        meas = it.get("measurements") or []
        if not meas:
            ws.append([idx, it.get("item_code"), it.get("description"), "Not measured", "", it.get("unit"),
                       0, 0, 0, 0, 0, rate, 0, "", "pending"])
            continue
        for m in meas:
            amt = round((m.get("measured_qty") or 0) * rate, 2)
            total += amt
            ws.append([idx, it.get("item_code"), it.get("description"), m.get("room_name"), m.get("floor"),
                       it.get("unit"), m.get("nos"), m.get("length"), m.get("breadth"), m.get("height"),
                       m.get("measured_qty"), rate, amt, fmt_date(m.get("measured_on")), m.get("status")])
    ws.append([])
    ws.append(["", "", "GRAND TOTAL", "", "", "", "", "", "", "", "", "", round(total, 2)])
    ws.cell(row=ws.max_row, column=3).font = bold
    ws.cell(row=ws.max_row, column=13).font = bold

    for row in ws.iter_rows(min_row=hrow, max_row=ws.max_row):
        for cell in row:
            cell.border = border
            if cell.column in (8, 9, 10, 11, 12, 13):
                cell.number_format = "#,##0.00"
            cell.alignment = Alignment(vertical="top", wrap_text=(cell.column == 3))
    widths = [8, 11, 58, 22, 14, 8, 7, 10, 10, 10, 11, 11, 14, 13, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ---- deviations sheet
    ws2 = wb.create_sheet("Deviation Statement")
    ws2.append(["CSR Code", "Short Description", "Unit", "Rate", "Tendered Qty", "Measured Qty", "Deviation",
                "Deviation %", "Excess/Saving Amount", "Alert"])
    for c in range(1, 11):
        cell = ws2.cell(row=1, column=c)
        cell.fill = hdr_fill
        cell.font = Font(bold=True, color="FFFFFF", size=9)
        cell.border = border
    for it in items:
        tq = it.get("tendered_qty") or 0
        mq = it.get("measured_qty") or 0
        rate = it.get("rate") or 0
        dev = round(mq - tq, 3)
        pct = (dev / tq * 100) if tq else 0
        alert = ""
        if tq and abs(pct) >= 20:
            alert = "CRITICAL - variation approval required"
        elif tq and abs(pct) >= 10:
            alert = "REVIEW - beyond permitted variation"
        elif dev > 0.001:
            alert = "Excess"
        elif dev < -0.001:
            alert = "Saving"
        ws2.append([it.get("item_code"), (it.get("short_desc") or it.get("description", ""))[:80], it.get("unit"),
                    rate, tq, mq, dev, round(pct, 2), round(dev * rate, 2), alert])
    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row):
        for cell in row:
            cell.border = border
            if cell.column in (4, 5, 6, 7, 9):
                cell.number_format = "#,##0.00"
            if cell.column == 10 and cell.value and "CRITICAL" in str(cell.value):
                cell.font = Font(bold=True, color="B91C1C")
            elif cell.column == 10 and cell.value and "REVIEW" in str(cell.value):
                cell.font = Font(color="B45309")
    for i, w in enumerate([11, 60, 8, 12, 13, 13, 11, 12, 20, 34], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # ---- project items sheet
    ws3 = wb.create_sheet("Estimate Items")
    ws3.append(["Item No.", "CSR Code", "Description (Master CSR)", "Unit", "Rate", "Tendered Qty",
                "Tendered Amount", "Source", "Non-Schedule", "Match Method", "Confidence"])
    for c in range(1, 12):
        cell = ws3.cell(row=1, column=c)
        cell.fill = hdr_fill
        cell.font = Font(bold=True, color="FFFFFF", size=9)
        cell.border = border
    for i, it in enumerate(items, start=1):
        ws3.append([i, it.get("item_code"), it.get("description"), it.get("unit"), it.get("rate"),
                    it.get("tendered_qty"), round((it.get("tendered_qty") or 0) * (it.get("rate") or 0), 2),
                    it.get("source"), "YES" if it.get("is_non_schedule") else "", it.get("match_method"),
                    it.get("confidence")])
    for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row):
        for cell in row:
            cell.border = border
            if cell.column in (5, 6, 7):
                cell.number_format = "#,##0.00"
    for i, w in enumerate([8, 11, 70, 8, 12, 13, 16, 10, 13, 18, 11], start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
